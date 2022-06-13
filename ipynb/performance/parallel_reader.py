from concurrent.futures import ProcessPoolExecutor
from openpyxl import load_workbook
from time import perf_counter

test_file = "Issues/bug494.xlsx"


def parallel_worksheet(sheetname):
    begin = perf_counter()
    wb = load_workbook(test_file, read_only=True,
                       data_only=True, keep_links=False)
    ws = wb[sheetname]
    for row in ws.iter_rows(values_only=True):
        for value in row:
            pass
    end = perf_counter()
    print("    {0} {1:.2f}s".format(sheetname, end - begin))


def parallel_read():
    print("Parallised Read")
    begin = perf_counter()
    wb = load_workbook(test_file, read_only=True,
                       keep_links=False, data_only=True)
    print("    Workbook loaded {0:.2f}s".format(perf_counter() - begin))
    sheets = wb.sheetnames
    with ProcessPoolExecutor() as pool:
        for ws in pool.map(parallel_worksheet, sheets):
            pass
    end = perf_counter()
    print("    Total time {0:.2f}s".format(end - begin))


if __name__ == "__main__":
    parallel_read()