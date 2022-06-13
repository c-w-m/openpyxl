"""
Compare read performance with xlrd and different openpyxl options
"""

try:
    from time import perf_counter
except ImportError:
    from time import clock as perf_counter
import sys

test_file = "Issues/bug494.xlsx"
import openpyxl
from openpyxl import load_workbook
import xlrd
from xlrd import open_workbook


def pixel(values_only=False, read_only=True):
    module_name = "openpyxl"
    if read_only:
        module_name += ", read-only"
    if values_only:
        module_name += ", values only"
    print(module_name)
    begin = perf_counter()

    wb = load_workbook(test_file, read_only=read_only, keep_links=False, data_only=True)
    print("    Workbook loaded {0:.2f}s".format(perf_counter() - begin))
    for ws in wb:
        start = perf_counter()
        rows = ws.iter_rows()
        if values_only:
            rows = ws.values
        for row in rows:
            for value in row:
                value
        end = perf_counter()
        print("    {0} {1:.2f}s".format(ws.title, end - start))
    print("    Total time {0:.2f}s".format(end - begin))
    print()


def xlread():
    print("xlrd")
    begin = perf_counter()

    wb = open_workbook(test_file)
    print("    Workbook loaded {0:.2f}s".format(perf_counter() - begin))
    for ws in wb.sheets():
        start = perf_counter()
        for idx in range(ws.nrows):
            for value in ws.row_values(idx):
                value
        end = perf_counter()
        print("    {0} {1:.2f}s".format(ws.name, end - start))
    print("    Total time {0:.2f}s".format(end - begin))
    print()


if __name__ == "__main__":
    print("")
    print("Versions:")
    print("%s: %s" % ('python', sys.version[:5]))
    print("%s: %s" % ('xlread', xlrd.__VERSION__))
    print("%s: %s" % ('openpyxl', openpyxl.__version__))
    print("")
    xlread()
    pixel(read_only=False)
    pixel(read_only=True)
    pixel(read_only=True, values_only=True)